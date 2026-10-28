"""Operaciones de caja y reportes de venta autorizados por rol y sede."""

from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
from typing import Literal

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import func
from sqlalchemy.orm import Session, joinedload

from app.auth.router import (
    SesionAutenticada,
    authorized_site_ids,
    get_current_session,
    require_cajero,
    require_selected_site,
)
from app.auditoria.models import EventoAuditoria
from app.core.clock import BUSINESS_TIMEZONE, get_utc_now
from app.core.database import get_db
from app.sedes.models import Sede
from app.usuarios.models import Usuario
from app.ventas.models import LineaPedido, Pedido
from app.ventas.schemas import (
    FacturaLineaOut,
    FacturaOut,
    PagoCreate,
    PedidoOut,
    PedidoPendienteOut,
    ReporteVentaOut,
)
from app.ventas.router import serialize_order

router = APIRouter(tags=["Caja y reportes"])


def build_invoice(pedido: Pedido) -> FacturaOut:
    """Construye el comprobante desde los valores históricos del pedido."""
    if pedido.cerrado_en is None or pedido.cerrado_por is None or pedido.medio_pago is None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="El pedido no tiene un pago registrado",
        )
    products = [
        FacturaLineaOut(
            producto_id=line.producto_id,
            codigo=line.codigo_producto,
            nombre=line.nombre_producto,
            cantidad=line.cantidad,
            precio_unitario=line.precio_unitario,
            subtotal=line.precio_unitario * line.cantidad,
        )
        for line in pedido.lineas
    ]
    return FacturaOut(
        pedido_id=pedido.id,
        sede_id=pedido.sede_id,
        mesa_id=pedido.mesa_id,
        creado_en=pedido.creado_en,
        cerrado_en=pedido.cerrado_en.astimezone(BUSINESS_TIMEZONE),
        cerrado_por=pedido.cerrado_por,
        medio_pago=pedido.medio_pago,
        productos=products,
        total=sum(
            (line.subtotal for line in products),
            start=Decimal("0.00"),
        ),
    )


@router.get(
    "/caja/pedidos-pendientes",
    response_model=list[PedidoPendienteOut],
)
def listar_pedidos_pendientes(
    sede: Sede = Depends(require_selected_site),
    _: Usuario = Depends(require_cajero),
    db: Session = Depends(get_db),
) -> list[PedidoPendienteOut]:
    """Lista pedidos abiertos de la sede seleccionada para el cajero."""
    orders = (
        db.query(Pedido)
        .options(joinedload(Pedido.lineas))
        .filter(Pedido.sede_id == sede.id, Pedido.estado == "ABIERTO")
        .order_by(Pedido.creado_en, Pedido.id)
        .all()
    )
    user_names = dict(
        db.query(Usuario.id, Usuario.nombre)
        .filter(Usuario.id.in_({order.usuario_id for order in orders}))
        .all()
    )
    return [
        PedidoPendienteOut(
            id=order.id,
            usuario_id=order.usuario_id,
            usuario_nombre=user_names[order.usuario_id],
            sede_id=order.sede_id,
            mesa_id=order.mesa_id,
            creado_en=order.creado_en,
            total=sum(
                (line.precio_unitario * line.cantidad for line in order.lineas),
                start=Decimal("0.00"),
            ),
        )
        for order in orders
    ]


@router.get("/caja/pedidos/{pedido_id}", response_model=PedidoOut)
def consultar_pedido_para_cobro(
    pedido_id: int,
    sede: Sede = Depends(require_selected_site),
    _: Usuario = Depends(require_cajero),
    db: Session = Depends(get_db),
) -> PedidoOut:
    """Devuelve detalle de un pedido de la sede activa al cajero."""
    pedido = (
        db.query(Pedido)
        .options(joinedload(Pedido.lineas))
        .filter(Pedido.id == pedido_id, Pedido.sede_id == sede.id)
        .first()
    )
    if pedido is None:
        raise HTTPException(status_code=404, detail="Pedido no encontrado")
    return serialize_order(pedido)


@router.post(
    "/caja/pedidos/{pedido_id}/pago",
    response_model=FacturaOut,
)
def registrar_pago(
    pedido_id: int,
    pago: PagoCreate,
    sede: Sede = Depends(require_selected_site),
    session: SesionAutenticada = Depends(get_current_session),
    _: Usuario = Depends(require_cajero),
    db: Session = Depends(get_db),
    now: datetime = Depends(get_utc_now),
) -> FacturaOut:
    """Registra el pago completo y cierra el pedido en una transacción."""
    pedido = (
        db.query(Pedido)
        .filter(Pedido.id == pedido_id, Pedido.sede_id == sede.id)
        .with_for_update()
        .first()
    )
    if pedido is None:
        raise HTTPException(status_code=404, detail="Pedido no encontrado")
    if pedido.estado != "ABIERTO":
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="El pedido ya fue cerrado",
        )
    pedido.estado = "CERRADO"
    pedido.medio_pago = pago.medio_pago
    pedido.cerrado_por = session.usuario.id
    pedido.cerrado_en = now
    db.add(
        EventoAuditoria(
            actor_id=session.usuario.id,
            sede_id=sede.id,
            pedido_id=pedido.id,
            accion="PEDIDO_CERRADO",
            creado_en=now,
        )
    )
    db.commit()
    db.refresh(pedido)
    return build_invoice(pedido)


@router.get("/caja/pedidos/{pedido_id}/factura", response_model=FacturaOut)
def consultar_factura(
    pedido_id: int,
    sede: Sede = Depends(require_selected_site),
    _: Usuario = Depends(require_cajero),
    db: Session = Depends(get_db),
) -> FacturaOut:
    """Muestra el comprobante de un pedido cerrado en la sede autorizada."""
    pedido = (
        db.query(Pedido)
        .options(joinedload(Pedido.lineas))
        .filter(
            Pedido.id == pedido_id,
            Pedido.sede_id == sede.id,
            Pedido.estado == "CERRADO",
        )
        .first()
    )
    if pedido is None:
        raise HTTPException(status_code=404, detail="Factura no encontrada")
    return build_invoice(pedido)


def query_sales_report(
    db: Session,
    user: Usuario,
    fecha_fin: date | None,
    sede_id: int | None,
    medio_pago: Literal["EFECTIVO", "TARJETA"] | None,
) -> list[ReporteVentaOut]:
    site_ids = authorized_site_ids(db, user)
    if sede_id is not None and sede_id not in site_ids:
        raise HTTPException(status_code=403, detail="No tienes acceso a esta sede")
    report_sites = [sede_id] if sede_id is not None else site_ids
    if not report_sites:
        return []

    query = (
        db.query(
            LineaPedido.producto_id.label("producto_id"),
            LineaPedido.codigo_producto.label("codigo"),
            LineaPedido.nombre_producto.label("nombre"),
            func.sum(LineaPedido.cantidad).label("cantidad"),
            func.sum(
                LineaPedido.cantidad * LineaPedido.precio_compra_unitario
            ).label("valor_compra"),
            LineaPedido.precio_unitario.label("precio_venta"),
            func.sum(
                LineaPedido.cantidad * LineaPedido.precio_unitario
            ).label("valor_venta"),
            Pedido.sede_id.label("sede_id"),
            Sede.nombre.label("sede"),
        )
        .join(Pedido, Pedido.id == LineaPedido.pedido_id)
        .join(Sede, Sede.id == Pedido.sede_id)
        .filter(Pedido.sede_id.in_(report_sites))
        .group_by(
            LineaPedido.producto_id,
            LineaPedido.codigo_producto,
            LineaPedido.nombre_producto,
            LineaPedido.precio_unitario,
            Pedido.sede_id,
            Sede.nombre,
        )
        .order_by(Sede.nombre, LineaPedido.nombre_producto)
    )
    query = query.filter(Pedido.estado == "CERRADO")
    if medio_pago is not None:
        query = query.filter(Pedido.medio_pago == medio_pago)
    if fecha_fin is not None:
        start = datetime.combine(fecha_fin, time.min, tzinfo=BUSINESS_TIMEZONE)
        end = datetime.combine(
            fecha_fin + timedelta(days=1),
            time.min,
            tzinfo=BUSINESS_TIMEZONE,
        )
        query = query.filter(Pedido.cerrado_en >= start, Pedido.cerrado_en < end)

    return [
        ReporteVentaOut(
            **{
                key: row._mapping[key]
                for key in (
                    "producto_id",
                    "codigo",
                    "nombre",
                    "cantidad",
                    "valor_compra",
                    "precio_venta",
                    "valor_venta",
                    "sede_id",
                    "sede",
                )
            },
            ganancia=row._mapping["valor_venta"] - row._mapping["valor_compra"],
        )
        for row in query.all()
    ]


@router.get("/reportes/ventas", response_model=list[ReporteVentaOut])
def consultar_reporte_ventas(
    fecha_fin: date | None = Query(default=None),
    sede_id: int | None = Query(default=None, gt=0),
    medio_pago: Literal["EFECTIVO", "TARJETA"] | None = None,
    session: SesionAutenticada = Depends(get_current_session),
    _: Usuario = Depends(require_cajero),
    db: Session = Depends(get_db),
) -> list[ReporteVentaOut]:
    """Filtra ventas sin permitir consultar sedes ajenas."""
    rows = query_sales_report(
        db,
        session.usuario,
        fecha_fin,
        sede_id,
        medio_pago,
    )
    db.add(
        EventoAuditoria(
            actor_id=session.usuario.id,
            sede_id=sede_id,
            accion="REPORTE_VENTAS_CONSULTADO",
            creado_en=get_utc_now(),
        )
    )
    db.commit()
    return rows


@router.get("/reportes/ventas/excel")
def exportar_reporte_ventas(
    fecha_fin: date | None = Query(default=None),
    sede_id: int | None = Query(default=None, gt=0),
    medio_pago: Literal["EFECTIVO", "TARJETA"] | None = None,
    session: SesionAutenticada = Depends(get_current_session),
    _: Usuario = Depends(require_cajero),
    db: Session = Depends(get_db),
) -> StreamingResponse:
    """Exporta exactamente las filas autorizadas y filtradas del reporte."""
    rows = query_sales_report(
        db,
        session.usuario,
        fecha_fin,
        sede_id,
        medio_pago,
    )
    db.add(
        EventoAuditoria(
            actor_id=session.usuario.id,
            sede_id=sede_id,
            accion="REPORTE_VENTAS_EXPORTADO",
            creado_en=get_utc_now(),
        )
    )
    db.commit()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ventas"
    sheet.append(
        [
            "ID producto",
            "Código",
            "Producto",
            "Cantidad",
            "Valor compra",
            "Precio venta unitario",
            "Valor venta",
            "Ganancia",
            "ID sede",
            "Sede",
        ]
    )
    for row in rows:
        sheet.append(
            [
                row.producto_id,
                row.codigo,
                row.nombre,
                row.cantidad,
                row.valor_compra,
                row.precio_venta,
                row.valor_venta,
                row.ganancia,
                row.sede_id,
                row.sede,
            ]
        )
    for column_cells in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(width + 2, 40)

    brown_fill = PatternFill(fill_type="solid", fgColor="38291A")
    cream_fill = PatternFill(fill_type="solid", fgColor="F5F0E5")
    alternate_cream_fill = PatternFill(fill_type="solid", fgColor="EEE3C8")
    gold_side = Side(style="thin", color="C6A45E")
    for cell in sheet[1]:
        cell.fill = brown_fill
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFF8E9")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=Side(style="medium", color="CFAD62"))
    sheet.row_dimensions[1].height = 30

    for row_index in range(2, sheet.max_row + 1):
        fill = cream_fill if row_index % 2 == 0 else alternate_cream_fill
        for cell in sheet[row_index]:
            cell.fill = fill
            cell.font = Font(name="Calibri", size=10, color="382D1C")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=gold_side)
        for column_index in (5, 6, 7, 8):
            sheet.cell(row=row_index, column=column_index).number_format = '"$"#,##0.00'

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = "CFAD62"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="reporte_ventas.xlsx"'},
    )
